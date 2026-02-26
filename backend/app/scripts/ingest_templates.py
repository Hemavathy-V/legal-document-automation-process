import argparse
import os
import re
import sys
from pathlib import Path

import mysql.connector
from docx import Document

# Add backend to path for logger import
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', '..', '..'))
from backend.app.core.logger import get_logger

logger = get_logger(__name__)

CLAUSE_TITLE_RE = re.compile(r"^(\d+)\.\s*(.+)$")
PLACEHOLDER_RE = re.compile(r"\{\{\s*([\w.]+)\s*\}\}")


def load_docx_text(path: Path) -> str:
    logger.debug(f"Loading DOCX text from: {path}")
    doc = Document(path)
    lines = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    result = "\n".join(lines)
    logger.debug(f"Loaded {len(lines)} lines from {path.name}")
    return result


def infer_template_name(path: Path) -> str:
    logger.debug(f"Inferring template name from: {path.name}")
    name = path.stem
    name = name.replace("-", " ").replace("_", " ")
    name = name.replace("Template", "").strip()
    result = " ".join(word.capitalize() for word in name.split())
    logger.debug(f"Inferred template name: {result}")
    return result


def extract_clauses(template_text: str):
    logger.debug("Extracting clauses from template")
    clauses = []
    current = None

    for line in template_text.splitlines():
        match = CLAUSE_TITLE_RE.match(line)
        if match:
            if current:
                clauses.append(current)
            current = {
                "number": int(match.group(1)),
                "title": match.group(2).strip(),
                "content": "",
            }
            continue

        if current:
            current["content"] += line + "\n"

    if current:
        clauses.append(current)

    logger.info(f"Extracted {len(clauses)} clauses from template")
    return clauses


def extract_placeholders(content: str):
    """Extract all unique placeholders from content."""
    logger.debug("Extracting placeholders from content")
    placeholders = set()
    for match in PLACEHOLDER_RE.finditer(content):
        placeholder = match.group(1)
        # Skip Handlebars block helpers like {{#services}}, {{/services}}
        if not placeholder.startswith('#') and not placeholder.startswith('/'):
            placeholders.add(placeholder)
    logger.debug(f"Found {len(placeholders)} unique placeholders")
    return sorted(placeholders)


def clause_flags(content: str):
    contains_placeholders = bool(PLACEHOLDER_RE.search(content))
    is_repeatable = "{{#" in content
    return contains_placeholders, is_repeatable


def get_connection():
    return mysql.connector.connect(
        host=os.environ.get("MYSQL_HOST", "localhost"),
        port=int(os.environ.get("MYSQL_PORT", "3306")),
        user=os.environ.get("MYSQL_USER", "root"),
        password=os.environ.get("MYSQL_PASSWORD", "root"),
        database=os.environ.get("MYSQL_DB", "legal_doc_template_db"),
    )


def insert_template(cursor, template_name, contract_type, version, template_content):
    logger.debug(f"Checking if template exists: {template_name} (v{version})")
    cursor.execute(
        """
        SELECT id FROM contract_templates
        WHERE template_name = %s AND version = %s
        """,
        (template_name, version),
    )
    row = cursor.fetchone()
    if row:
        logger.debug(f"Template already exists: {template_name} (ID: {row[0]})")
        return row[0], False

    logger.debug(f"Inserting new template: {template_name}")
    cursor.execute(
        """
        INSERT INTO contract_templates
            (template_name, contract_type, version, template_content)
        VALUES (%s, %s, %s, %s)
        """,
        (template_name, contract_type, version, template_content),
    )
    template_id = cursor.lastrowid
    logger.info(f"Template inserted: {template_name} (ID: {template_id})")
    return template_id, True


def insert_clauses(cursor, template_id, clauses):
    logger.debug(f"Inserting {len(clauses)} clauses for template ID: {template_id}")
    for clause in clauses:
        contains_placeholders, is_repeatable = clause_flags(clause["content"])
        placeholders = extract_placeholders(clause["content"])
        # Store placeholders as comma-separated string in clause_title
        clause_title_value = ", ".join(placeholders) if placeholders else clause["title"]
        
        logger.debug(f"Inserting clause {clause['number']}: {clause['title']}")
        cursor.execute(
            """
            INSERT INTO template_clauses
                (template_id, clause_number, clause_title, is_mandatory,
                 is_repeatable, contains_placeholders)
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (
                template_id,
                clause["number"],
                clause_title_value,
                True,
                is_repeatable,
                contains_placeholders,
            ),
        )
    logger.info(f"Inserted all {len(clauses)} clauses for template ID: {template_id}")


def main():
    logger.info("Starting template ingestion process")
    
    parser = argparse.ArgumentParser(description="Ingest DOCX templates into MySQL")
    parser.add_argument(
        "--templates-dir",
        default=str(Path(__file__).resolve().parents[1] / "sample_templates"),
    )
    parser.add_argument("--version", default="v1.0")
    args = parser.parse_args()
    
    logger.debug(f"Templates directory: {args.templates_dir}")
    logger.debug(f"Template version: {args.version}")

    templates_dir = Path(args.templates_dir)
    if not templates_dir.exists():
        logger.error(f"Templates directory not found: {templates_dir}")
        raise SystemExit(f"Templates dir not found: {templates_dir}")

    # Filter out temporary Word files (starting with ~$)
    docx_files = sorted(
        f for f in templates_dir.glob("*.docx") 
        if not f.name.startswith("~$")
    )
    
    logger.info(f"Found {len(docx_files)} template files in {templates_dir}")
    
    if not docx_files:
        logger.error("No DOCX templates found")
        raise SystemExit("No .docx templates found")

    connection = get_connection()
    cursor = connection.cursor()
    
    logger.info("Database connection established")

    successfully_processed = 0
    skipped = 0
    failed = 0
    
    for path in docx_files:
        try:
            logger.info(f"Processing template: {path.name}")
            template_content = load_docx_text(path)
            template_name = infer_template_name(path)
            contract_type = template_name

            template_id, created = insert_template(
                cursor, template_name, contract_type, args.version, template_content
            )

            if created:
                clauses = extract_clauses(template_content)
                if clauses:
                    insert_clauses(cursor, template_id, clauses)
                successfully_processed += 1
                logger.info(f"Inserted: {template_name} ({path.name})")
            else:
                skipped += 1
                logger.info(f"Skipped: {template_name} ({path.name}) - already exists")
        except Exception as e:
            failed += 1
            logger.error(f"Error processing {path.name}: {str(e)}")
            continue

    connection.commit()
    cursor.close()
    connection.close()
    
    logger.info(f"Template ingestion completed: {successfully_processed} inserted, {skipped} skipped, {failed} failed")
    logger.info("Database connection closed")


if __name__ == "__main__":
    main()