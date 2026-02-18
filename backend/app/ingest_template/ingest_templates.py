import argparse
import os
import re
from pathlib import Path

import mysql.connector
from docx import Document

CLAUSE_TITLE_RE = re.compile(r"^(\d+)\.\s*(.+)$")
PLACEHOLDER_RE = re.compile(r"\{\{\s*([\w.]+)\s*\}\}")


def load_docx_text(path: Path) -> str:
    doc = Document(path)
    lines = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)
    return "\n".join(lines)


def infer_template_name(path: Path) -> str:
    name = path.stem
    name = name.replace("-", " ").replace("_", " ")
    name = name.replace("Template", "").strip()
    return " ".join(word.capitalize() for word in name.split())


def extract_clauses(template_text: str):
    clauses = []
    current = None

    for line in template_text.splitlines():
        match = CLAUSE_TITLE_RE.match(line)
        if match:
            if current:
                clauses.append(current)
            current = {
                "number": int(match.group(1)),
                "title": match.group(2).strip(),
                "content": "",
            }
            continue

        if current:
            current["content"] += line + "\n"

    if current:
        clauses.append(current)

    return clauses


def extract_placeholders(content: str):
    """Extract all unique placeholders from content."""
    placeholders = set()
    for match in PLACEHOLDER_RE.finditer(content):
        placeholder = match.group(1)
        # Skip Handlebars block helpers like {{#services}}, {{/services}}
        if not placeholder.startswith('#') and not placeholder.startswith('/'):
            placeholders.add(placeholder)
    return sorted(placeholders)


def clause_flags(content: str):
    contains_placeholders = bool(PLACEHOLDER_RE.search(content))
    is_repeatable = "{{#" in content
    return contains_placeholders, is_repeatable


def get_connection():
    return mysql.connector.connect(
        host=os.environ.get("MYSQL_HOST", "localhost"),
        port=int(os.environ.get("MYSQL_PORT", "3306")),
        user=os.environ.get("MYSQL_USER", "root"),
        password=os.environ.get("MYSQL_PASSWORD", "root"),
        database=os.environ.get("MYSQL_DB", "legal_doc_template_db"),
    )


def insert_template(cursor, template_name, contract_type, version, template_content):
    cursor.execute(
        """
        SELECT id FROM contract_templates
        WHERE template_name = %s AND version = %s
        """,
        (template_name, version),
    )
    row = cursor.fetchone()
    if row:
        return row[0], False

    cursor.execute(
        """
        INSERT INTO contract_templates
            (template_name, contract_type, version, template_content)
        VALUES (%s, %s, %s, %s)
        """,
        (template_name, contract_type, version, template_content),
    )
    return cursor.lastrowid, True


def insert_clauses(cursor, template_id, clauses):
    for clause in clauses:
        contains_placeholders, is_repeatable = clause_flags(clause["content"])
        placeholders = extract_placeholders(clause["content"])
        # Store placeholders as comma-separated string in clause_title
        clause_title_value = ", ".join(placeholders) if placeholders else clause["title"]
        
        cursor.execute(
            """
            INSERT INTO template_clauses
                (template_id, clause_number, clause_title, is_mandatory,
                 is_repeatable, contains_placeholders)
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (
                template_id,
                clause["number"],
                clause_title_value,
                True,
                is_repeatable,
                contains_placeholders,
            ),
        )


def main():
    parser = argparse.ArgumentParser(description="Ingest DOCX templates into MySQL")
    parser.add_argument(
        "--templates-dir",
        default=str(Path(__file__).resolve().parents[1] / "sample_templates"),
    )
    parser.add_argument("--version", default="v1.0")
    args = parser.parse_args()

    templates_dir = Path(args.templates_dir)
    if not templates_dir.exists():
        raise SystemExit(f"Templates dir not found: {templates_dir}")

    # Filter out temporary Word files (starting with ~$)
    docx_files = sorted(
        f for f in templates_dir.glob("*.docx") 
        if not f.name.startswith("~$")
    )
    if not docx_files:
        raise SystemExit("No .docx templates found")

    connection = get_connection()
    cursor = connection.cursor()

    for path in docx_files:
        try:
            template_content = load_docx_text(path)
            template_name = infer_template_name(path)
            contract_type = template_name

            template_id, created = insert_template(
                cursor, template_name, contract_type, args.version, template_content
            )

            if created:
                clauses = extract_clauses(template_content)
                if clauses:
                    insert_clauses(cursor, template_id, clauses)

            action = "Inserted" if created else "Skipped"
            print(f"{action}: {template_name} ({path.name})")
        except Exception as e:
            print(f"Error processing {path.name}: {e}")
            continue

    connection.commit()
    cursor.close()
    connection.close()


if __name__ == "__main__":
    main()
