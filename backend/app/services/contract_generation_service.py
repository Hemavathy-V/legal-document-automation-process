import json
from pathlib import Path
from typing import Dict

from docx import Document

from backend.app.core.gemini_client import get_gemini_model
from backend.app.core.logger import get_logger
from backend.app.prompts.legal_contract_prompt import CONTRACT_PROMPT
from backend.app.services.ollama_service import generate_response as generate_ollama_response


logger = get_logger(__name__)


class ContractGenerationService:
    def __init__(self, output_dir: Path):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.model = get_gemini_model()

    def load_json(self, json_path: Path) -> Dict:
        with open(json_path, "r", encoding="utf-8") as file_obj:
            return json.load(file_obj)

    def build_prompt(
        self,
        contract_type: str,
        template_text: str,
        user_input_data: dict,
        regional_law_text: str,
        kb_clauses_text: str,
    ) -> str:
        return CONTRACT_PROMPT.format(
            contract_type=contract_type,
            contract_template_text=template_text,
            user_input_json=json.dumps(user_input_data, indent=2, ensure_ascii=False),
            regional_law_text=regional_law_text,
            kb_clauses_text=kb_clauses_text,
        )

    def _validate_generated_contract(self, generated_text: str) -> str:
        if not generated_text or not generated_text.strip():
            raise ValueError("Model returned empty response")

        cleaned_text = generated_text.strip()
        if "{{" in cleaned_text or "}}" in cleaned_text:
            raise ValueError("Generated contract still contains unresolved placeholders")

        if len(cleaned_text) < 200:
            raise ValueError("Generated contract response is too short to be valid")

        return cleaned_text

    def _is_quota_or_rate_limit_error(self, error: Exception) -> bool:
        message = str(error).lower()
        return (
            "429" in message
            or "quota" in message
            or "rate limit" in message
            or "retry in" in message
        )

    def _generate_with_ollama(self, prompt: str) -> str:
        logger.info("Falling back to Ollama for contract generation")
        response_text = generate_ollama_response(prompt)

        if not response_text:
            raise RuntimeError("Ollama returned empty response")

        lowered = response_text.lower()
        if (
            "ollama server is not running" in lowered
            or "model timeout" in lowered
            or lowered.startswith("unexpected error:")
        ):
            raise RuntimeError(f"Ollama generation failed: {response_text}")

        return self._validate_generated_contract(response_text)

    def generate_contract(self, prompt: str) -> str:
        logger.debug(f"Prompt length: {len(prompt)}")

        last_error = None
        for attempt in range(2):
            try:
                logger.info(f"Calling Gemini for contract generation (attempt {attempt + 1}/2)")
                response = self.model.generate_content(
                    prompt,
                    generation_config={
                        "temperature": 0.2,
                        "max_output_tokens": 4096,
                    },
                )

                response_text = getattr(response, "text", None)
                if not response_text and getattr(response, "candidates", None):
                    parts = getattr(response.candidates[0].content, "parts", [])
                    response_text = "".join(getattr(part, "text", "") for part in parts)

                return self._validate_generated_contract(response_text or "")
            except Exception as exc:
                last_error = exc
                logger.warning(f"Gemini generation attempt failed: {str(exc)}")

                if self._is_quota_or_rate_limit_error(exc):
                    logger.warning("Gemini quota/rate limit detected; switching to Ollama fallback")
                    try:
                        return self._generate_with_ollama(prompt)
                    except Exception as ollama_exc:
                        raise RuntimeError(
                            f"Gemini quota exceeded and Ollama fallback failed: {ollama_exc}"
                        ) from ollama_exc

        try:
            return self._generate_with_ollama(prompt)
        except Exception as ollama_exc:
            raise RuntimeError(
                f"Gemini generation failed after retries: {last_error}; "
                f"Ollama fallback failed: {ollama_exc}"
            ) from ollama_exc

    def save_generated_contract(self, contract_id: int, text: str) -> Path:
        filename = f"generated_contract_{contract_id}.docx"
        path = self.output_dir / filename

        document = Document()
        for block in text.split("\n\n"):
            normalized = block.strip()
            if normalized:
                document.add_paragraph(normalized)

        document.save(path)
        logger.info(f"Generated contract DOCX saved at {path}")
        return path
